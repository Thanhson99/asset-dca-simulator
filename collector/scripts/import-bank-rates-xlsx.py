#!/usr/bin/env python3
"""Import historical bank deposit rates from the curated Excel workbook.

The workbook is treated as data input only. This script converts the
Historical_Obs sheet into the JSON snapshot format consumed by the app.
"""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BANKS = {
    "Vietcombank": {
        "provider": "VCB",
        "product_id": "VCB_COUNTER_SAVING",
        "name": "Vietcombank tiết kiệm VND",
    },
    "BIDV": {
        "provider": "BIDV",
        "product_id": "BIDV_COUNTER_SAVING",
        "name": "BIDV tiết kiệm VND",
    },
    "VietinBank": {
        "provider": "CTG",
        "product_id": "CTG_COUNTER_SAVING",
        "name": "VietinBank tiết kiệm VND",
    },
    "Agribank": {
        "provider": "AGRIBANK",
        "product_id": "AGRIBANK_COUNTER_SAVING",
        "name": "Agribank tiết kiệm VND",
    },
    "Techcombank": {
        "provider": "TCB",
        "product_id": "TCB_COUNTER_SAVING",
        "name": "Techcombank tiết kiệm VND",
    },
    "MB": {
        "provider": "MBB",
        "product_id": "MBB_COUNTER_SAVING",
        "name": "MB tiết kiệm VND",
    },
    "VPBank": {
        "provider": "VPB",
        "product_id": "VPB_COUNTER_SAVING",
        "name": "VPBank tiết kiệm VND",
    },
    "ACB": {
        "provider": "ACB",
        "product_id": "ACB_COUNTER_SAVING",
        "name": "ACB tiết kiệm VND",
    },
    "Sacombank": {
        "provider": "STB",
        "product_id": "STB_COUNTER_SAVING",
        "name": "Sacombank tiết kiệm VND",
    },
    "TPBank": {
        "provider": "TPB",
        "product_id": "TPB_COUNTER_SAVING",
        "name": "TPBank tiết kiệm VND",
    },
}

TERM_COLUMNS = ["1T", "2T", "3T", "4T", "6T", "9T", "12T", "13T", "15T", "18T", "24T", "36T", "48T", "60T"]


def main() -> None:
    parser = argparse.ArgumentParser(description="Import bank deposit rates from an Excel workbook.")
    parser.add_argument("workbook", help="Path to lich_su_lai_suat_10_ngan_hang_2000_2026.xlsx")
    parser.add_argument("--repo-root", default=".", help="Repository root. Defaults to current directory.")
    args = parser.parse_args()

    repo_root = Path(args.repo_root).resolve()
    workbook_path = Path(args.workbook).expanduser().resolve()
    wb = load_workbook(workbook_path, read_only=True, data_only=True)

    historical_rows = read_historical_rows(wb)
    annual_rows = read_coverage_rows(wb, "Annual_Snapshot")
    coverage_rows = read_coverage_rows(wb, "Coverage")
    source_rows = read_coverage_rows(wb, "Sources")

    histories = build_histories(historical_rows, annual_rows, coverage_rows, source_rows, workbook_path)
    for provider, history in histories.items():
      output = repo_root / "data" / "rates" / "banks" / provider / "history.json"
      output.parent.mkdir(parents=True, exist_ok=True)
      write_json(output, history)
      print(f"[OK] {provider}: {len(history['snapshots'])} snapshot(s) -> {output.relative_to(repo_root)}")

    update_index(repo_root, histories)
    print(f"Imported {sum(len(item['snapshots']) for item in histories.values())} bank rate snapshot(s).")


def read_historical_rows(wb: Any) -> list[dict[str, Any]]:
    ws = wb["Historical_Obs"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(value).strip() for value in rows[0]]
    data = []

    for row in rows[1:]:
        values = dict(zip(headers, row))
        if not values.get("Ngân hàng") or not values.get("Ngày") or not values.get("Kỳ hạn (tháng)"):
            continue

        bank_name = str(values["Ngân hàng"]).strip()
        if bank_name not in BANKS:
            raise ValueError(f"Unsupported bank name: {bank_name}")

        rate = number_or_none(values.get("Lãi suất (%/năm)"))
        if rate is None:
            continue

        data.append(
            {
                "bank": bank_name,
                "date": iso_date(values["Ngày"]),
                "term": int(values["Kỳ hạn (tháng)"]),
                "rate": rate,
                "channel": clean_text(values.get("Kênh")),
                "product": clean_text(values.get("Sản phẩm")),
                "condition": clean_text(values.get("Điều kiện")),
                "confidence": clean_text(values.get("Độ tin cậy")),
                "note": clean_text(values.get("Ghi chú")),
                "issuer": clean_text(values.get("Đơn vị/Nguồn phát hành")),
                "source": clean_text(values.get("Tên nguồn")) or clean_text(values.get("Đơn vị/Nguồn phát hành")) or "Excel import",
                "sourceUrl": clean_text(values.get("URL")),
            }
        )

    return data


def read_coverage_rows(wb: Any, sheet_name: str) -> list[dict[str, Any]]:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    data = []

    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue
        values = dict(zip(headers, row))
        data.append({key: normalize_metadata_value(value) for key, value in values.items() if key})

    return data


def build_histories(
    rows: list[dict[str, Any]],
    annual_rows: list[dict[str, Any]],
    coverage_rows: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    workbook_path: Path,
) -> dict[str, dict[str, Any]]:
    rows_by_provider = defaultdict(list)
    for row in rows:
        rows_by_provider[BANKS[row["bank"]]["provider"]].append(row)

    histories = {}
    generated_at = datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    for bank_name, config in BANKS.items():
        provider = config["provider"]
        bank_rows = sorted(rows_by_provider[provider], key=lambda item: (item["term"], item["date"], item["channel"], item["condition"]))
        effective_to_by_key = compute_effective_to(bank_rows)
        snapshots = []

        for row in sorted(bank_rows, key=lambda item: (item["date"], item["term"], item["channel"], item["condition"])):
            key = rate_key(row)
            source = "xlsx-import"
            if row["source"]:
                source = f"xlsx-import: {row['source']}"

            snapshots.append(
                {
                    "date": row["date"],
                    "effectiveFrom": row["date"],
                    "effectiveTo": effective_to_by_key.get(key),
                    "channel": row["channel"] or "unknown",
                    "product": row["product"] or "Tiết kiệm VND KHCN - lãi cuối kỳ",
                    "condition": row["condition"] or "Tiêu chuẩn",
                    "confidence": row["confidence"] or None,
                    "source": source,
                    "sourceUrl": row["sourceUrl"] or "",
                    "sourceName": row["source"] or None,
                    "sourceIssuer": row["issuer"] or None,
                    "note": row["note"] or None,
                    "needsReview": (row["confidence"] or "").upper() != "A",
                    "terms": [{"months": row["term"], "annualRate": row["rate"]}],
                }
            )

        histories[provider] = compact_none(
            {
                "version": 1,
                "productId": config["product_id"],
                "type": "bank_saving",
                "provider": provider,
                "name": config["name"],
                "currency": "VND",
                "generatedAt": generated_at,
                "sourceWorkbook": workbook_path.name,
                "sourceWorkbookSheets": ["Historical_Obs", "Annual_Snapshot", "Coverage", "Sources"],
                "coverage": find_bank_metadata(coverage_rows, bank_name),
                "annualSnapshots": [item for item in annual_rows if item.get("Ngân hàng") == bank_name],
                "sources": source_rows,
                "snapshots": snapshots,
            }
        )

    return histories


def compute_effective_to(rows: list[dict[str, Any]]) -> dict[tuple[Any, ...], str | None]:
    by_term = defaultdict(list)
    for row in rows:
        by_term[row["term"]].append(row)

    effective_to = {}
    for term_rows in by_term.values():
        dates = sorted({row["date"] for row in term_rows})
        next_date_by_date = {
            current_date: dates[index + 1] if index + 1 < len(dates) else None
            for index, current_date in enumerate(dates)
        }
        for row in term_rows:
            next_date = next_date_by_date[row["date"]]
            effective_to[rate_key(row)] = previous_day(next_date) if next_date else None

    return effective_to


def update_index(repo_root: Path, histories: dict[str, dict[str, Any]]) -> None:
    index_path = repo_root / "data" / "rates" / "index.json"
    with index_path.open("r", encoding="utf-8") as file:
        index = json.load(file)

    by_provider = {history["provider"]: history for history in histories.values()}
    for product in index["products"]:
        history = by_provider.get(product.get("provider"))
        if not history:
            continue

        terms = sorted({term["months"] for snapshot in history["snapshots"] for term in snapshot["terms"]})
        product.update(
            {
                "id": history["productId"],
                "name": history["name"],
                "terms": terms,
                "dataPath": f"data/rates/banks/{history['provider']}/history.json",
            }
        )

    write_json(index_path, index)


def find_bank_metadata(rows: list[dict[str, Any]], bank_name: str) -> dict[str, Any] | None:
    for row in rows:
        if row.get("Ngân hàng") == bank_name:
            return row
    return None


def rate_key(row: dict[str, Any]) -> tuple[Any, ...]:
    return (row["bank"], row["date"], row["term"], row["channel"], row["condition"], row["sourceUrl"])


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def number_or_none(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def iso_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return datetime.fromisoformat(str(value)).date().isoformat()


def previous_day(value: str) -> str:
    return (datetime.fromisoformat(value).date() - timedelta(days=1)).isoformat()


def normalize_metadata_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def compact_none(value: Any) -> Any:
    if isinstance(value, dict):
        return {key: compact_none(item) for key, item in value.items() if item is not None}
    if isinstance(value, list):
        return [compact_none(item) for item in value]
    return value


def write_json(path: Path, data: Any) -> None:
    with path.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)
        file.write("\n")


if __name__ == "__main__":
    main()
